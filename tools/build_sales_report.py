#!/usr/bin/env python3
"""
月次売上分析レポート生成パイプライン

使い方:
    python3 build_sales_report.py --raw <生データxlsx> [--previous <前月の出力xlsx>] [--out-dir <出力先ディレクトリ>]

仕様の詳細は docs/売上分析_出力仕様.md を参照。
このスクリプトには特定の月名・年月・目標金額・店舗名・商品名・施策名を一切
ハードコードしていない。すべて --raw / --previous で渡されるファイルと
report_config.json（表記ゆれのうち機械的正規化・読み仮名一致で解決できない
同義語だけを載せる、人手で保守する小さな対応表）から読み取る。
"""
import argparse
import json
import re
import unicodedata
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import pykakasi

# ---------------------------------------------------------------------------
# 固定スキーマ定数（データ値ではなく「表の構造」を表す定数のみ。
#  月・年・金額・店舗名・商品名・施策名などのデータ値は一切含まない）
# ---------------------------------------------------------------------------
FONT_NAME = "Arial"
BODY_SIZE = 11
HEADER_FILL = "1F3864"
HEADER_FONT_COLOR = "FFFFFF"
YEN_FMT = "¥#,##0"
PCT_FMT = "0.0%"
DATE_FMT = "yyyy/m/d"

SHEET_ORDER = ["サマリー", "日次推移", "店舗別", "チャネル別", "カテゴリ別", "顧客区分別", "クリーニング済明細"]

REQUIRED_RAW_COLS = ["注文ID", "注文日", "店舗", "チャネル", "顧客名", "顧客区分",
                      "カテゴリ", "商品名", "数量", "単価", "割引率", "支払方法", "備考"]

DETAIL_COLUMNS = REQUIRED_RAW_COLS + ["有効フラグ", "単価補完", "売上金額", "原価単価", "粗利"]

DIM_HEADERS = ["項目", "件数", "売上金額", "粗利", "粗利率", "新規項目"]
DAILY_HEADERS = ["日付", "件数", "売上金額", "粗利"]

# 列幅は「列の内容」に対して固定（月によって変えない）
COLUMN_WIDTHS = {
    "注文ID": 14, "注文日": 12, "日付": 12, "店舗": 14, "チャネル": 12,
    "顧客名": 16, "顧客区分": 12, "カテゴリ": 12, "商品名": 22, "数量": 8,
    "単価": 12, "割引率": 10, "支払方法": 14, "備考": 18, "有効フラグ": 10,
    "単価補完": 10, "売上金額": 14, "原価単価": 12, "粗利": 14, "粗利率": 10,
    "件数": 10, "項目": 14, "新規項目": 10,
}

# 店舗名などの地名系の表記ゆれを読み仮名で突き合わせる際に取り除く、
# 業種一般に使われる接尾語（特定の店舗名ではない）
GENERIC_SUFFIXES = {
    "店舗": ["店舗", "支店", "店"],
    "顧客区分": ["顧客", "客"],
    "チャネル": [],
    "カテゴリ": [],
}

ZEN2HAN_TABLE = str.maketrans(
    {chr(0xFF01 + i): chr(0x21 + i) for i in range(0x5E)} | {"　": " "}
)

_KKS = pykakasi.kakasi()


def zenkaku_to_hankaku(s):
    return s.translate(ZEN2HAN_TABLE)


def mech_key(s):
    """空白除去＋全角/半角統一＋（ASCIIのみの場合）大文字化による機械的な正規化キー"""
    if s is None:
        return ""
    s = zenkaku_to_hankaku(str(s))
    s = re.sub(r"\s+", "", s)
    if s.isascii():
        s = s.upper()
    return s


def reading_key(s, axis):
    core = re.sub(r"\s+", "", zenkaku_to_hankaku(str(s)))
    for suf in GENERIC_SUFFIXES.get(axis, []):
        if core.endswith(suf) and len(core) > len(suf):
            core = core[: -len(suf)]
            break
    conv = _KKS.convert(core)
    return "".join(x["hira"] for x in conv)


def load_alias_config(path):
    if path is None or not Path(path).exists():
        return {}
    with open(path, encoding="utf-8") as f:
        cfg = json.load(f)
    cfg.pop("_comment", None)
    return cfg


# ---------------------------------------------------------------------------
# 軸（店舗・チャネル・顧客区分・カテゴリ）の正規化と表示順の決定
# ---------------------------------------------------------------------------
class AxisResolver:
    """生データ中の表記ゆれを正規名に解決し、表示順（前月ファイル優先、
    無ければ当月データの初出順）を決める。新規に追加された正規名は
    new_items に記録する。"""

    def __init__(self, axis, alias_map, previous_order):
        self.axis = axis
        self.alias_map = alias_map or {}
        self.previous_order = list(previous_order) if previous_order else []
        self._canonical_by_mech = {mech_key(c): c for c in self.previous_order}
        self._canonical_by_reading = {reading_key(c, axis): c for c in self.previous_order}
        self.order = list(self.previous_order)
        self.new_items = []
        self._seen_this_run = OrderedDict()

    def resolve(self, raw_value):
        raw = "" if raw_value is None else str(raw_value).strip()
        if raw == "":
            raw = "（区分なし）" if self.axis == "顧客区分" else raw
        if raw in self.alias_map:
            raw = self.alias_map[raw]

        mkey = mech_key(raw)
        if mkey in self._canonical_by_mech:
            canonical = self._canonical_by_mech[mkey]
        else:
            rkey = reading_key(raw, self.axis)
            if rkey and rkey in self._canonical_by_reading:
                canonical = self._canonical_by_reading[rkey]
            else:
                canonical = raw
                self._canonical_by_mech[mkey] = canonical
                if rkey:
                    self._canonical_by_reading[rkey] = canonical

        if canonical not in self._seen_this_run:
            self._seen_this_run[canonical] = True
            if canonical not in self.order:
                self.order.append(canonical)
                self.new_items.append(canonical)
        return canonical

    def final_order(self):
        # previous_order の並びを保持しつつ、当月データに一度も出現しなかった
        # 前月項目も 0 件として残す（呼び出し側で order をそのまま使えばよい）
        return self.order


# ---------------------------------------------------------------------------
# 生データ読み込み・構造クレンジング
# ---------------------------------------------------------------------------
DATA_SHEET_RE = re.compile(r"^\d{1,2}月売上$")


def find_data_sheet(wb):
    for name in wb.sheetnames:
        if DATA_SHEET_RE.match(name):
            return name
    raise ValueError("売上データのシート（『<月>月売上』という名前）が見つかりません")


def find_header_row(ws, max_scan=15):
    for r in range(1, max_scan + 1):
        if ws.cell(row=r, column=1).value == "注文ID":
            return r
    raise ValueError("ヘッダー行（A列が『注文ID』の行）が見つかりません")


def is_blank_row(row):
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def load_raw_records(raw_path):
    wb = openpyxl.load_workbook(raw_path, data_only=True)
    sheet_name = find_data_sheet(wb)
    ws = wb[sheet_name]
    header_row = find_header_row(ws)
    header = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
    col_idx = {h: i for i, h in enumerate(header)}
    for req in REQUIRED_RAW_COLS:
        if req not in col_idx:
            raise ValueError(f"必須列『{req}』がヘッダー行に見つかりません")

    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, len(header) + 1)]
        if is_blank_row(row):
            continue
        oid = row[col_idx["注文ID"]]
        date_val = row[col_idx["注文日"]]
        product_val = row[col_idx["商品名"]]

        def _blank(v):
            return v is None or (isinstance(v, str) and not v.strip())

        if _blank(oid) or _blank(date_val) or _blank(product_val):
            continue  # タイトル行・合計行・メモ行など（注文ID/注文日/商品名が欠ける行）
        records.append(row)

    seen = set()
    deduped = []
    for row in records:
        oid = row[col_idx["注文ID"]]
        if oid in seen:
            continue
        seen.add(oid)
        deduped.append(row)

    if "商品マスタ" not in wb.sheetnames:
        raise ValueError("『商品マスタ』シートが見つかりません")
    if "メモ" not in wb.sheetnames:
        raise ValueError("『メモ』シートが見つかりません")

    master_ws = wb["商品マスタ"]
    product_category = {}
    product_price = {}
    product_cost = {}
    for row in master_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name, category, std_price, cost = row[0], row[1], row[2], row[3]
        product_category[name] = category
        product_price[name] = std_price
        product_cost[name] = cost
    category_order = list(dict.fromkeys(product_category.values()))

    memo_ws = wb["メモ"]
    memo_lines = [row[0] for row in memo_ws.iter_rows(values_only=True) if row and row[0]]

    return {
        "col_idx": col_idx,
        "rows": deduped,
        "product_category": product_category,
        "product_price": product_price,
        "product_cost": product_cost,
        "category_order": category_order,
        "memo_lines": memo_lines,
    }


# ---------------------------------------------------------------------------
# 型の正規化（日付・数量・単価・割引率）
# ---------------------------------------------------------------------------
import datetime  # noqa: E402

CANCEL_RETURN_KEYWORDS = ["キャンセル", "返品"]


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return datetime.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return datetime.datetime.strptime(str(int(v)), "%Y%m%d")
    if isinstance(v, str):
        s = zenkaku_to_hankaku(v.strip())
        m = re.match(r"^(\d{1,2})月(\d{1,2})日$", s)
        if m:
            return ("month_day", int(m.group(1)), int(m.group(2)))
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.datetime.strptime(s, fmt)
            except ValueError:
                pass
    raise ValueError(f"注文日を解釈できません: {v!r}")


def parse_qty(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v)
    s = zenkaku_to_hankaku(str(v)).strip()
    s = re.sub(r"[^\d]", "", s)
    return int(s)


def parse_price(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v)
    s = zenkaku_to_hankaku(str(v)).strip()
    s = re.sub(r"[^\d]", "", s)
    return int(s) if s else None


def parse_discount(v):
    if isinstance(v, str):
        s = zenkaku_to_hankaku(v.strip()).rstrip("%")
        return float(s) / 100
    v = float(v)
    return v / 100 if v > 1 else v


def is_cancel_or_return(note):
    return isinstance(note, str) and any(k in note for k in CANCEL_RETURN_KEYWORDS)


def build_clean_detail(loaded, resolvers):
    """loaded: load_raw_records() の戻り値, resolvers: axis -> AxisResolver
    戻り値: (detail_rows[dict], report_year, report_month)"""
    col_idx = loaded["col_idx"]
    product_category = loaded["product_category"]
    product_price = loaded["product_price"]
    product_cost = loaded["product_cost"]

    detail_rows = []
    month_counter = Counter()
    fallback_year_month = None

    # 先に年月を仮決定するため、日付だけ一巡してモードを取る
    parsed_dates = []
    for row in loaded["rows"]:
        d = parse_date(row[col_idx["注文日"]])
        parsed_dates.append(d)
        if isinstance(d, datetime.datetime):
            month_counter[(d.year, d.month)] += 1
    if month_counter:
        report_year, report_month = month_counter.most_common(1)[0][0]
    else:
        raise ValueError("注文日から年月を特定できませんでした（絶対年月を含む日付が1件もありません）")

    for row, d in zip(loaded["rows"], parsed_dates):
        if isinstance(d, tuple):  # ("month_day", m, day) 形式 -> レポート年を採用
            _, m, day = d
            d = datetime.datetime(report_year, m, day)

        product = row[col_idx["商品名"]]
        if product not in product_category:
            raise ValueError(f"商品マスタに存在しない商品名です: {product!r}")

        price = parse_price(row[col_idx["単価"]])
        filled = False
        if price is None:
            price = product_price[product]
            filled = True

        qty = parse_qty(row[col_idx["数量"]])
        discount = parse_discount(row[col_idx["割引率"]])
        note = row[col_idx["備考"]]
        valid = "無効" if is_cancel_or_return(note) else "有効"

        revenue = qty * price * (1 - discount)
        cost = product_cost[product]
        profit = revenue - qty * cost

        detail = {
            "注文ID": row[col_idx["注文ID"]],
            "注文日": d,
            "店舗": resolvers["店舗"].resolve(row[col_idx["店舗"]]),
            "チャネル": resolvers["チャネル"].resolve(row[col_idx["チャネル"]]),
            "顧客名": row[col_idx["顧客名"]],
            "顧客区分": resolvers["顧客区分"].resolve(row[col_idx["顧客区分"]]),
            "カテゴリ": resolvers["カテゴリ"].resolve(product_category[product]),
            "商品名": product,
            "数量": qty,
            "単価": price,
            "割引率": discount,
            "支払方法": row[col_idx["支払方法"]],
            "備考": note,
            "有効フラグ": valid,
            "単価補完": "補完" if filled else "",
            "売上金額": revenue,
            "原価単価": cost,
            "粗利": profit,
        }
        detail_rows.append(detail)

    return detail_rows, report_year, report_month


# ---------------------------------------------------------------------------
# メモシートの読み取り（目標額・施策名・施策対象期間）
# 記載が見つからない場合は None を返す（推測しない）
# ---------------------------------------------------------------------------
def parse_memo(memo_lines, report_year):
    text = "\n".join(str(x) for x in memo_lines)

    target = None
    for line in text.split("\n"):
        if "目標" not in line:
            continue
        m = re.search(r"([\d,]{4,})\s*円", line)
        if m:
            target = int(m.group(1).replace(",", ""))
            break

    campaign_name = None
    campaign_start = None
    campaign_end = None
    for line in text.split("\n"):
        qm = re.search(r"[「『]([^」』]+)[」』]", line)
        dm = re.search(r"(\d{1,2})/(\d{1,2})\s*[-~〜]\s*(\d{1,2})/(\d{1,2})", line)
        if qm and dm:
            campaign_name = qm.group(1)
            campaign_start = datetime.datetime(report_year, int(dm.group(1)), int(dm.group(2)))
            campaign_end = datetime.datetime(report_year, int(dm.group(3)), int(dm.group(4)))
            break

    return {
        "target_amount": target,
        "campaign_name": campaign_name,
        "campaign_start": campaign_start,
        "campaign_end": campaign_end,
    }


# ---------------------------------------------------------------------------
# 出力ワークブックの組み立て
# ---------------------------------------------------------------------------
def _style_header_cell(cell):
    cell.font = Font(name=FONT_NAME, size=BODY_SIZE, bold=True, color=HEADER_FONT_COLOR)
    cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_body_cell(cell, bold=False):
    cell.font = Font(name=FONT_NAME, size=BODY_SIZE, bold=bold)


def _apply_column_widths(ws, headers):
    for i, h in enumerate(headers, start=1):
        width = COLUMN_WIDTHS.get(h, 14)
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_header_row(ws, headers, row=1):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        _style_header_cell(cell)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def build_detail_sheet(wb, detail_rows):
    ws = wb.create_sheet("クリーニング済明細")
    _write_header_row(ws, DETAIL_COLUMNS)
    for r, d in enumerate(detail_rows, start=2):
        for c, h in enumerate(DETAIL_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=d[h])
            _style_body_cell(cell)
            if h in ("注文日",):
                cell.number_format = DATE_FMT
            elif h in ("単価", "売上金額", "原価単価", "粗利"):
                cell.number_format = YEN_FMT
            elif h == "割引率":
                cell.number_format = PCT_FMT
    _apply_column_widths(ws, DETAIL_COLUMNS)
    return ws


def build_dim_sheet(wb, name, axis_key, order, new_items, detail_rows, show_new_flag):
    """項目ごとに 件数・売上金額・粗利・粗利率 を集計するシートを作る。
    order にある項目は当月の実績がゼロでも行を残す。new_items は当月新たに
    追加された項目（前月ファイルとの比較で判明したもの）。"""
    ws = wb.create_sheet(name)
    _write_header_row(ws, DIM_HEADERS)
    valid_rows = [d for d in detail_rows if d["有効フラグ"] == "有効"]
    new_set = set(new_items) if show_new_flag else set()

    r = 2
    for item in order:
        rows_for_item = [d for d in valid_rows if d[axis_key] == item]
        count = len(rows_for_item)
        revenue = sum(d["売上金額"] for d in rows_for_item)
        profit = sum(d["粗利"] for d in rows_for_item)
        margin = (profit / revenue) if revenue else 0

        ws.cell(row=r, column=1, value=item)
        ws.cell(row=r, column=2, value=count)
        ws.cell(row=r, column=3, value=revenue)
        ws.cell(row=r, column=4, value=profit)
        ws.cell(row=r, column=5, value=margin)
        ws.cell(row=r, column=6, value=("新規項目" if item in new_set else ""))
        for c in range(1, 7):
            _style_body_cell(ws.cell(row=r, column=c))
        ws.cell(row=r, column=3).number_format = YEN_FMT
        ws.cell(row=r, column=4).number_format = YEN_FMT
        ws.cell(row=r, column=5).number_format = PCT_FMT
        r += 1

    total_row = r
    total_count = sum(1 for d in valid_rows)
    total_revenue = sum(d["売上金額"] for d in valid_rows)
    total_profit = sum(d["粗利"] for d in valid_rows)
    total_margin = (total_profit / total_revenue) if total_revenue else 0
    ws.cell(row=total_row, column=1, value="合計")
    ws.cell(row=total_row, column=2, value=total_count)
    ws.cell(row=total_row, column=3, value=total_revenue)
    ws.cell(row=total_row, column=4, value=total_profit)
    ws.cell(row=total_row, column=5, value=total_margin)
    for c in range(1, 6):
        _style_body_cell(ws.cell(row=total_row, column=c), bold=True)
    ws.cell(row=total_row, column=3).number_format = YEN_FMT
    ws.cell(row=total_row, column=4).number_format = YEN_FMT
    ws.cell(row=total_row, column=5).number_format = PCT_FMT

    _apply_column_widths(ws, DIM_HEADERS)
    return ws, total_row


def build_daily_sheet(wb, report_year, report_month, detail_rows, campaign_start, campaign_end):
    import calendar

    ws = wb.create_sheet("日次推移")
    headers = DAILY_HEADERS + ["施策期間"]
    _write_header_row(ws, headers)

    valid_rows = [d for d in detail_rows if d["有効フラグ"] == "有効"]
    by_date = {}
    for d in valid_rows:
        key = d["注文日"].date()
        agg = by_date.setdefault(key, {"count": 0, "revenue": 0, "profit": 0})
        agg["count"] += 1
        agg["revenue"] += d["売上金額"]
        agg["profit"] += d["粗利"]

    n_days = calendar.monthrange(report_year, report_month)[1]
    r = 2
    for day in range(1, n_days + 1):
        cur = datetime.date(report_year, report_month, day)
        agg = by_date.get(cur, {"count": 0, "revenue": 0, "profit": 0})
        in_campaign = ""
        if campaign_start and campaign_end:
            if campaign_start.date() <= cur <= campaign_end.date():
                in_campaign = "期間内"
        ws.cell(row=r, column=1, value=datetime.datetime(cur.year, cur.month, cur.day))
        ws.cell(row=r, column=2, value=agg["count"])
        ws.cell(row=r, column=3, value=agg["revenue"])
        ws.cell(row=r, column=4, value=agg["profit"])
        ws.cell(row=r, column=5, value=in_campaign)
        for c in range(1, 6):
            _style_body_cell(ws.cell(row=r, column=c))
        ws.cell(row=r, column=1).number_format = DATE_FMT
        ws.cell(row=r, column=3).number_format = YEN_FMT
        ws.cell(row=r, column=4).number_format = YEN_FMT
        r += 1

    total_row = r
    ws.cell(row=total_row, column=1, value="合計")
    ws.cell(row=total_row, column=2, value=sum(d["count"] for d in by_date.values()))
    ws.cell(row=total_row, column=3, value=sum(d["revenue"] for d in by_date.values()))
    ws.cell(row=total_row, column=4, value=sum(d["profit"] for d in by_date.values()))
    for c in range(1, 5):
        _style_body_cell(ws.cell(row=total_row, column=c), bold=True)
    ws.cell(row=total_row, column=3).number_format = YEN_FMT
    ws.cell(row=total_row, column=4).number_format = YEN_FMT

    _apply_column_widths(ws, headers)
    return ws, n_days, total_row


def build_summary_sheet(wb, report_year, report_month, detail_rows, memo_info,
                         daily_sheet_info, store_sheet_info, all_new_items, previous_given):
    ws = wb.create_sheet("サマリー", 0)

    title = ws.cell(row=1, column=1, value=f"{report_year}年{report_month}月 売上サマリー")
    title.font = Font(name=FONT_NAME, size=16, bold=True)

    valid_rows = [d for d in detail_rows if d["有効フラグ"] == "有効"]
    revenue = sum(d["売上金額"] for d in valid_rows)
    profit = sum(d["粗利"] for d in valid_rows)
    margin = (profit / revenue) if revenue else 0
    order_count = len(valid_rows)
    aov = (revenue / order_count) if order_count else 0
    target = memo_info["target_amount"]
    achievement = (revenue / target) if target else None

    kpi_labels = ["売上高", "粗利", "粗利率", "注文件数", "客単価", "目標達成率"]
    kpi_values = [revenue, profit, margin, order_count, aov,
                  achievement if achievement is not None else "記載なし"]
    kpi_formats = [YEN_FMT, YEN_FMT, PCT_FMT, "#,##0", YEN_FMT, PCT_FMT]

    header_row, value_row = 3, 4
    for i, (label, value, fmt) in enumerate(zip(kpi_labels, kpi_values, kpi_formats)):
        col = 2 + i  # B..G
        hcell = ws.cell(row=header_row, column=col, value=label)
        _style_header_cell(hcell)
        vcell = ws.cell(row=value_row, column=col, value=value)
        _style_body_cell(vcell, bold=True)
        if isinstance(value, (int, float)):
            vcell.number_format = fmt
        ws.column_dimensions[get_column_letter(col)].width = 14

    info_row = 6
    ws.cell(row=info_row, column=1, value="月間売上目標")
    ws.cell(row=info_row, column=2,
            value=target if target is not None else "記載なし")
    if target is not None:
        ws.cell(row=info_row, column=2).number_format = YEN_FMT
    ws.cell(row=info_row + 1, column=1, value="販促施策名")
    ws.cell(row=info_row + 1, column=2, value=memo_info["campaign_name"] or "記載なし")
    ws.cell(row=info_row + 2, column=1, value="施策対象期間")
    if memo_info["campaign_start"] and memo_info["campaign_end"]:
        period_text = (f"{memo_info['campaign_start'].strftime('%Y/%m/%d')} - "
                        f"{memo_info['campaign_end'].strftime('%Y/%m/%d')}")
    else:
        period_text = "記載なし"
    ws.cell(row=info_row + 2, column=2, value=period_text)
    for rr in range(info_row, info_row + 3):
        _style_body_cell(ws.cell(row=rr, column=1), bold=True)
        _style_body_cell(ws.cell(row=rr, column=2))

    new_row = info_row + 4
    ws.cell(row=new_row, column=1, value="新規項目")
    if not previous_given:
        note = "初回実行のため比較対象となる前月ファイルがなく、新規項目の判定は行っていません（本ファイルが以後の基準になります）。"
    elif all_new_items:
        note = "、".join(f"{axis}:{name}" for axis, name in all_new_items)
    else:
        note = "なし"
    ws.cell(row=new_row, column=2, value=note)
    _style_body_cell(ws.cell(row=new_row, column=1), bold=True)
    _style_body_cell(ws.cell(row=new_row, column=2))
    ws.merge_cells(start_row=new_row, start_column=2, end_row=new_row, end_column=8)

    ws.column_dimensions["A"].width = 16

    # --- 折れ線グラフ（日次推移） ---
    daily_ws, n_days, daily_total_row = daily_sheet_info
    line = LineChart()
    line.title = "日次売上推移"
    line.style = 2
    line.y_axis.title = "売上金額（円）"
    line.x_axis.title = "日付"
    line.width, line.height = 18, 9
    data_ref = Reference(daily_ws, min_col=3, max_col=3, min_row=1, max_row=daily_total_row - 1)
    cats_ref = Reference(daily_ws, min_col=1, max_col=1, min_row=2, max_row=daily_total_row - 1)
    line.add_data(data_ref, titles_from_data=True)
    line.set_categories(cats_ref)
    ws.add_chart(line, "A14")

    # --- 棒グラフ（店舗別） ---
    store_ws, store_total_row = store_sheet_info
    bar = BarChart()
    bar.type = "col"
    bar.title = "店舗別売上"
    bar.style = 10
    bar.y_axis.title = "売上金額（円）"
    bar.x_axis.title = "店舗"
    bar.width, bar.height = 18, 9
    bdata_ref = Reference(store_ws, min_col=3, max_col=3, min_row=1, max_row=store_total_row - 1)
    bcats_ref = Reference(store_ws, min_col=1, max_col=1, min_row=2, max_row=store_total_row - 1)
    bar.add_data(bdata_ref, titles_from_data=True)
    bar.set_categories(bcats_ref)
    ws.add_chart(bar, "M14")

    return ws


# ---------------------------------------------------------------------------
# 前月出力ファイルから軸の並び順を読み取る
# ---------------------------------------------------------------------------
SHEET_TO_AXIS = {"店舗別": "店舗", "チャネル別": "チャネル", "カテゴリ別": "カテゴリ", "顧客区分別": "顧客区分"}


def load_previous_orders(previous_path):
    if not previous_path:
        return {}
    wb = openpyxl.load_workbook(previous_path, data_only=True)
    orders = {}
    for sheet_name, axis in SHEET_TO_AXIS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        items = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[0]
            if name is None or name == "合計":
                continue
            items.append(name)
        orders[axis] = items
    return orders


# ---------------------------------------------------------------------------
# メイン処理
# ---------------------------------------------------------------------------
def build_report(raw_path, previous_path=None, config_path=None, out_dir=None):
    if config_path is None:
        config_path = Path(__file__).parent / "report_config.json"
    alias_config = load_alias_config(config_path)
    previous_orders = load_previous_orders(previous_path)
    previous_given = bool(previous_path)

    loaded = load_raw_records(raw_path)

    resolvers = {
        "店舗": AxisResolver("店舗", alias_config.get("店舗"), previous_orders.get("店舗")),
        "チャネル": AxisResolver("チャネル", alias_config.get("チャネル"), previous_orders.get("チャネル")),
        "顧客区分": AxisResolver("顧客区分", alias_config.get("顧客区分"), previous_orders.get("顧客区分")),
        "カテゴリ": AxisResolver("カテゴリ", {}, previous_orders.get("カテゴリ") or loaded["category_order"]),
    }
    # 商品マスタに存在するのにまだ order に無いカテゴリを補っておく（商品マスタを基準とする仕様）
    for cat in loaded["category_order"]:
        if cat not in resolvers["カテゴリ"].order:
            resolvers["カテゴリ"].order.append(cat)

    detail_rows, report_year, report_month = build_clean_detail(loaded, resolvers)
    memo_info = parse_memo(loaded["memo_lines"], report_year)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_detail_sheet(wb, detail_rows)

    daily_ws, n_days, daily_total_row = build_daily_sheet(
        wb, report_year, report_month, detail_rows,
        memo_info["campaign_start"], memo_info["campaign_end"]
    )

    dim_specs = [
        ("店舗別", "店舗", resolvers["店舗"]),
        ("チャネル別", "チャネル", resolvers["チャネル"]),
        ("カテゴリ別", "カテゴリ", resolvers["カテゴリ"]),
        ("顧客区分別", "顧客区分", resolvers["顧客区分"]),
    ]
    dim_sheet_info = {}
    all_new_items = []
    for sheet_name, axis_key, resolver in dim_specs:
        ws, total_row = build_dim_sheet(
            wb, sheet_name, axis_key, resolver.final_order(), resolver.new_items,
            detail_rows, show_new_flag=previous_given
        )
        dim_sheet_info[sheet_name] = (ws, total_row)
        if previous_given:
            for item in resolver.new_items:
                all_new_items.append((axis_key, item))

    build_summary_sheet(
        wb, report_year, report_month, detail_rows, memo_info,
        (daily_ws, n_days, daily_total_row),
        dim_sheet_info["店舗別"],
        all_new_items, previous_given,
    )

    for name in SHEET_ORDER:
        wb.move_sheet(name, offset=(SHEET_ORDER.index(name) - wb.sheetnames.index(name)))

    out_dir = Path(out_dir) if out_dir else Path(raw_path).parent
    out_path = out_dir / f"{report_year}年{report_month}月_売上分析.xlsx"
    wb.save(out_path)
    return out_path


def main():
    parser = argparse.ArgumentParser(description="月次売上分析レポート生成")
    parser.add_argument("--raw", required=True, help="生データxlsxファイルのパス")
    parser.add_argument("--previous", help="前月の出力xlsxファイルのパス（軸の並び順・新規項目判定に使用）")
    parser.add_argument("--config", help="report_config.json のパス（省略時はスクリプトと同じ場所）")
    parser.add_argument("--out-dir", help="出力先ディレクトリ（省略時は --raw と同じ場所）")
    args = parser.parse_args()

    out_path = build_report(args.raw, args.previous, args.config, args.out_dir)
    print(f"saved: {out_path}")


if __name__ == "__main__":
    main()
